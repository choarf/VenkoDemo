"""Test and weekly report generation (HTML + XLSX).

Pure functions: callers fetch Athena/DynamoDB data and pass plain rows in, so
this module is unit-tested without AWS (backend/tests/test_logic.py).
"""
import io
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HIGHCHARTS_VERSION = "12.1.2"
_TEMPLATES = Path(__file__).resolve().parent.parent / "templates"
_env = Environment(loader=FileSystemLoader(_TEMPLATES), autoescape=select_autoescape(["html", "j2"]))


# ------------------------------------------------------------- analysis
def collapse_episodes(rows: list[dict], max_gap_ms: int) -> list[dict]:
    """Merge consecutive abnormal samples (same sensor + kind) into episodes.

    rows: {k, t, val, kind} sorted or not; kind is LOW/HIGH or a read error
    status (BUS_ERROR/EXCEPTION). A break longer than max_gap_ms starts a new
    episode, since the gap means samples in between were normal (or missing).
    """
    episodes = []
    cur = None
    for r in sorted(rows, key=lambda r: (r["k"], r["kind"], r["t"])):
        same = cur and cur["k"] == r["k"] and cur["kind"] == r["kind"] and r["t"] - cur["end"] <= max_gap_ms
        if same:
            cur["end"] = r["t"]
            cur["samples"] += 1
            cur["peak"] = _peak(cur["kind"], cur["peak"], r.get("val"))
        else:
            cur = {"k": r["k"], "kind": r["kind"], "start": r["t"], "end": r["t"], "samples": 1,
                   "peak": r.get("val")}
            episodes.append(cur)
    return sorted(episodes, key=lambda e: e["start"])


def _peak(kind, a, b):
    if a is None:
        return b
    if b is None:
        return a
    return min(a, b) if kind == "LOW" else max(a, b)


def _local(ms, tz, fmt="%Y-%m-%d %H:%M:%S"):
    return datetime.fromtimestamp(ms / 1000, ZoneInfo(tz)).strftime(fmt) if ms is not None else ""


def _r(v, nd=2):
    return None if v is None else round(float(v), nd)


def build_context(*, kind: str, title: str, config: dict, from_ms: int, to_ms: int,
                  summary: list[dict], series: list[dict], bucket_ms: int,
                  abnormal: list[dict], gaps: list[dict], system: dict | None,
                  comments: list[dict], meta: dict | None = None, generated_ms: int) -> dict:
    tz = config.get("gateway", {}).get("city") or "UTC"
    poll_ms = int(config.get("gateway", {}).get("poll_interval", 20)) * 1000
    sensors_cfg = {s["key"]: s for s in config.get("sensors", [])}
    by_key = {row["k"]: row for row in summary}
    expected = max(1, (to_ms - from_ms) // poll_ms)

    sensors = []
    for key, sc in sensors_cfg.items():
        row = by_key.get(key, {})
        ok = row.get("ok") or 0
        in_alarm = (row.get("low") or 0) + (row.get("high") or 0)
        status = "SIN DATOS" if not ok else ("ALARMA" if in_alarm else "OK")
        sensors.append({
            "key": key, "label": sc.get("label", key), "unit": sc.get("unit", ""),
            "lim_min": sc.get("min"), "lim_max": sc.get("max"),
            "samples": row.get("samples") or 0, "ok": ok,
            "availability": _r(min(100.0, 100.0 * ok / expected), 1),
            "min": _r(row.get("min")), "avg": _r(row.get("avg")), "max": _r(row.get("max")),
            "std": _r(row.get("std"), 3), "first": _r(row.get("first")), "last": _r(row.get("last")),
            "low": row.get("low") or 0, "high": row.get("high") or 0,
            "alarm_pct": _r(100.0 * in_alarm / ok, 1) if ok else None,
            "status": status,
        })

    chart_series = {}
    for p in series:
        chart_series.setdefault(p["k"], []).append([p["t"], _r(p["avg"], 3), _r(p["min"], 3), _r(p["max"], 3)])

    episodes = []
    for e in collapse_episodes(abnormal, max_gap_ms=int(poll_ms * 2.5)):
        sc = sensors_cfg.get(e["k"], {})
        episodes.append({
            **e, "label": sc.get("label", e["k"]), "unit": sc.get("unit", ""),
            "start_local": _local(e["start"], tz), "end_local": _local(e["end"], tz),
            "minutes": _r((e["end"] - e["start"] + poll_ms) / 60000, 1), "peak": _r(e["peak"]),
        })

    gap_rows = [{
        "start_local": _local(g["start"], tz), "end_local": _local(g["stop"], tz),
        "minutes": _r((g["stop"] - g["start"]) / 60000, 1),
    } for g in gaps]
    missing_ms = sum(g["stop"] - g["start"] - poll_ms for g in gaps)

    comment_rows = [{
        "time_local": _local(c["ts_ms"], tz), "author": c.get("author", ""),
        "sensor": sensors_cfg.get(c.get("sensor") or "", {}).get("label", c.get("sensor") or "General"),
        "text": c.get("text", ""),
    } for c in sorted(comments, key=lambda c: c["ts_ms"])]

    return {
        "kind": kind, "title": title, "tz": tz, "meta": meta or {},
        "site": config.get("site", {}), "gateway": config.get("gateway", {}),
        "from_ms": from_ms, "to_ms": to_ms,
        "from_local": _local(from_ms, tz), "to_local": _local(to_ms, tz),
        "duration_h": _r((to_ms - from_ms) / 3_600_000, 2),
        "generated_local": _local(generated_ms, tz),
        "sensors": sensors, "series": chart_series, "bucket_min": _r(bucket_ms / 60000, 2),
        "episodes": episodes, "gaps": gap_rows,
        "data_availability": _r(max(0.0, 100.0 * (1 - missing_ms / max(1, to_ms - from_ms))), 1),
        "system": system or {}, "comments": comment_rows,
        "alarm_sensor_count": sum(1 for s in sensors if s["status"] == "ALARMA"),
        "highcharts_version": HIGHCHARTS_VERSION,
    }


# ------------------------------------------------------------- rendering
def render_html(ctx: dict) -> bytes:
    tpl = _env.get_template("report.html.j2")
    # Chart data is embedded as JSON; escape '</' so it can't close the script tag.
    charts = json.dumps([
        {"key": s["key"], "label": s["label"], "unit": s["unit"], "min": s["lim_min"], "max": s["lim_max"],
         "data": ctx["series"].get(s["key"], [])}
        for s in ctx["sensors"]
    ]).replace("</", "<\\/")
    return tpl.render(ctx=ctx, charts_json=charts).encode("utf-8")


_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="0F2235")
_ALARM_FILL = PatternFill("solid", fgColor="FDE2E1")


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = _HEAD, _HEAD_FILL, Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, len(str(h)) + 2))
    return ws


def render_xlsx(ctx: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    info = [
        ("Reporte", ctx["title"]),
        ("Sitio", ctx["site"].get("title", "")),
        ("Gateway", ctx["gateway"].get("gateway_id", "")),
        ("Desde", ctx["from_local"]), ("Hasta", ctx["to_local"]),
        ("Duración (h)", ctx["duration_h"]), ("Zona horaria", ctx["tz"]),
        ("Disponibilidad de datos (%)", ctx["data_availability"]),
        ("Sensores en alarma", ctx["alarm_sensor_count"]),
        ("Generado", ctx["generated_local"]),
    ] + [(k.capitalize(), v) for k, v in ctx["meta"].items() if v not in (None, "")]
    for k, v in info:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    ws.append([])
    head = ["Sensor", "Clave", "Unidad", "Lím. mín", "Lím. máx", "Muestras", "Disp. %", "Mín", "Prom", "Máx",
            "Desv. est.", "Primero", "Último", "Bajo", "Alto", "% en alarma", "Estado"]
    start = ws.max_row + 1
    ws.append(head)
    for c in ws[start]:
        c.font, c.fill = _HEAD, _HEAD_FILL
    for s in ctx["sensors"]:
        ws.append([s["label"], s["key"], s["unit"], s["lim_min"], s["lim_max"], s["samples"], s["availability"],
                   s["min"], s["avg"], s["max"], s["std"], s["first"], s["last"], s["low"], s["high"],
                   s["alarm_pct"], s["status"]])
        if s["status"] != "OK":
            for c in ws[ws.max_row]:
                c.fill = _ALARM_FILL
    for i in range(3, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, 11)

    # Wide readings table: one row per time bucket, one column per sensor (average).
    keys = [s["key"] for s in ctx["sensors"]]
    grid: dict[int, dict] = {}
    for k, pts in ctx["series"].items():
        for t, avg, *_ in pts:
            grid.setdefault(t, {})[k] = avg
    labels = [f'{s["label"]} ({s["unit"]})' for s in ctx["sensors"]]
    _sheet(wb, f'Lecturas ({ctx["bucket_min"]} min)', ["Hora local"] + labels,
           [[_local(t, ctx["tz"])] + [grid[t].get(k) for k in keys] for t in sorted(grid)],
           widths={"Hora local": 20})

    _sheet(wb, "Alarmas", ["Sensor", "Tipo", "Inicio", "Fin", "Minutos", "Muestras", "Pico", "Unidad"],
           [[e["label"], e["kind"], e["start_local"], e["end_local"], e["minutes"], e["samples"], e["peak"], e["unit"]]
            for e in ctx["episodes"]], widths={"Sensor": 24, "Inicio": 20, "Fin": 20})
    _sheet(wb, "Huecos de datos", ["Inicio", "Fin", "Minutos"],
           [[g["start_local"], g["end_local"], g["minutes"]] for g in ctx["gaps"]],
           widths={"Inicio": 20, "Fin": 20})
    _sheet(wb, "Comentarios", ["Hora", "Autor", "Sensor", "Comentario"],
           [[c["time_local"], c["author"], c["sensor"], c["text"]] for c in ctx["comments"]],
           widths={"Hora": 20, "Sensor": 24, "Comentario": 80})
    sysd = ctx["system"]
    _sheet(wb, "Sistema", ["Métrica", "Valor"], [[k, v] for k, v in sysd.items()], widths={"Métrica": 20, "Valor": 40})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
