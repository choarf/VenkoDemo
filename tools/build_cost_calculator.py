"""Build docs/aws_cost_calculator.xlsx - formula-driven AWS cost model for one VenkoDemo site."""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F2F2F2")
NOTE = Font(name=F, italic=True, color="595959", size=9)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '$#,##0.00;($#,##0.00);"-"'
USD4 = '$#,##0.0000;($#,##0.0000);"-"'
USDP = '$#,##0.0000000000'
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM4 = '#,##0.0000;(#,##0.0000);"-"'

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def header(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR, HDR_FILL, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SEC_FILL
    ws.cell(row=row, column=1, value=text).font = BOLD


# ---------------------------------------------------------------- Summary (filled last, created first)
sm = wb.active
sm.title = "Summary"

# ---------------------------------------------------------------- Sensors
ss = wb.create_sheet("Sensors")
ss["A1"] = "Sensor list (one row per sensor published by the Akvo_Green gateway)"
ss["A1"].font = TITLE
ss["A2"] = ("Edit the blue cells: Device ID and Sensor name exactly as in Akvo_Green config.json, Active = 1 to include, 0 to exclude. "
            "Add sensors in the empty rows below (up to row 205). Pre-filled from Akvo_Green/config_data/config.json (15 sensors).")
ss["A2"].font = NOTE
ss["A3"] = "Name lengths matter: every reading repeats the device and sensor name in the MQTT JSON payload."
ss["A3"].font = NOTE
header(ss, 5, ["#", "Device ID", "Sensor name", "Active (1/0)", "Bytes per reading", "Device bytes (first sensor of device)"],
       [6, 16, 22, 12, 16, 22])
ss.row_dimensions[5].height = 30
sensors = [("DEV_1", "Pt100Sensor"), ("DEV_2", "Ambient"), ("DEV_3", "Humedad"), ("DEV_4", "ORPSensor"),
           ("DEV_5", "PhSensor"), ("DEV_6", "TempPhSensor"), ("DEV_7", "DOSensor"), ("DEV_8", "CODSensor"),
           ("DEV_9", "TempCODSensor"), ("DEV_10", "TurbCODSensort"), ("DEV_11", "BODSensor"), ("DEV_12", "TOCSensor"),
           ("DEV_13", "CondSensor"), ("DEV_14", "TempCondSensor"), ("DEV_15", "TDSCondSensor")]
FIRST, LAST = 6, 205
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=True,
                    error="Use 1 (active) or 0 (inactive)", showErrorMessage=True)
ss.add_data_validation(dv)
for r in range(FIRST, LAST + 1):
    i = r - FIRST
    ss.cell(row=r, column=1, value=i + 1).font = NOTE
    dev, sen = sensors[i] if i < len(sensors) else (None, None)
    for col, v in ((2, dev), (3, sen), (4, 1 if dev else None)):
        c = ss.cell(row=r, column=col, value=v)
        c.font, c.border = BLUE, BOX
    dv.add(f"D{r}")
    ss.cell(row=r, column=5, value=f'=IF(AND(D{r}=1,C{r}<>""),LEN(C{r})+sensor_over_b,0)').font = BLACK
    ss.cell(row=r, column=6,
            value=f'=IF(AND(E{r}>0,COUNTIFS(B${FIRST}:B{r},B{r},D${FIRST}:D{r},1,C${FIRST}:C{r},"<>")=1),LEN(B{r})+device_over_b,0)').font = BLACK
    for col in (5, 6):
        ss.cell(row=r, column=col).number_format = NUM
        ss.cell(row=r, column=col).border = BOX
ss.freeze_panes = "A6"
SENS_E = f"Sensors!$E${FIRST}:$E${LAST}"
SENS_F = f"Sensors!$F${FIRST}:$F${LAST}"

# ---------------------------------------------------------------- Inputs
ins = wb.create_sheet("Inputs")
ins["A1"] = "Inputs and assumptions (per site)"
ins["A1"].font = TITLE
ins["A2"] = "Blue = value you can change · yellow fill = main levers · grey section = technical constants derived from the code (change only if the code changes)."
ins["A2"].font = NOTE
header(ins, 4, ["Parameter", "Name", "Value", "Unit", "Source / note"], [46, 22, 14, 18, 90])
r = 5


def inp(label, nm, value, unit, note, key=False, tech=False, fmt=None):
    global r
    ins.cell(row=r, column=1, value=label).font = BLACK
    ins.cell(row=r, column=2, value=nm).font = NOTE
    c = ins.cell(row=r, column=3, value=value)
    c.font, c.border = BLUE, BOX
    if fmt:
        c.number_format = fmt
    if key:
        c.fill = YELLOW
    elif tech:
        c.fill = GREY
    ins.cell(row=r, column=4, value=unit).font = BLACK
    ins.cell(row=r, column=5, value=note).font = NOTE
    name(nm, f"Inputs!$C${r}")
    r += 1


def sec(text):
    global r
    r += 1
    section(ins, r, text, 5)
    r += 1


sec("Gateway (Akvo_Green config.json → gateway section)")
inp("Sensor data interval", "poll_s", 20, "seconds", "gateway.poll_interval — one data message with ALL sensors per interval", key=True)
inp("System data interval", "sys_s", 60, "seconds", "gateway.system_interval — CPU/RAM/disk/IP/OS message", key=True)
inp("Data topic", "topic_data", "AKVO/data", "text", "aws.topic_pub (stack DataTopic). Its length is part of every MQTT packet")
inp("System topic", "topic_sys", "AKVO/system", "text", "aws.topic_system (stack SystemTopic)")
inp("Number of sites (Pi + stack + dashboard)", "sites", 1, "sites", "Each site is a full separate stack; the total is multiplied by this", key=True)
inp("Days per month", "days_month", "=365/12", "days", "Average month")

sec("Dashboard usage")
inp("Dashboards open at the same time (average)", "viewers", 1, "browsers", "Each open dashboard polls /live + /tests/active", key=True)
inp("Hours per day a dashboard is open", "hours_open", 8, "hours/day", "Left open 24 h on a control-room screen = 24", key=True)
inp("Live poll interval", "live_poll_s", 10, "seconds", "LIVE_POLL_MS in web/js/main.js")
inp("Dashboard page loads", "page_loads", 10, "loads/day", "Each load downloads web/ through CloudFront (caching disabled) and calls config/active/lists")
inp("History / alarm queries", "history_q", 20, "queries/day", "Every History or Alarms view refresh runs one Athena query", key=True)
inp("Average query range", "history_days", 1, "days", "Range selected in History/Alarms (max 31 days, MAX_RANGE_DAYS)", key=True)
inp("Excel/CSV exports", "exports", 10, "exports/month", "GET /export runs one Athena query")

sec("Tests and reports")
inp("Tests per month", "tests_month", 8, "tests/month", "Each Start/Stop test produces one report (3 Athena queries)", key=True)
inp("Average test duration", "test_hours", 4, "hours", "Report queries scan whole days")
inp("Weekly reports per month", "weekly_reports", "=52/12", "reports/month", "EventBridge every Monday")

sec("Retention and pricing options")
inp("Months of history kept in S3 / CloudWatch", "months_kept", 12, "months", "Raw data never expires (bucket Retain, no expiry rule). Storage cost shown is the month after this many months", key=True)
inp("Apply AWS always-free tier", "free_tier", "No", "Yes/No", "Yes = subtract always-free allowances (Lambda, CloudFront, CloudWatch, DynamoDB storage, Glue, data transfer). One allowance per AWS account, shared by all sites", key=True)
free_row = r - 1

sec("Technical constants (from the code - change only if the code changes)")
inp("Data payload fixed part", "data_base_b", 57, "bytes", 'json.dumps({"ts": "<ISO time>", "devices": {}}) in edge_node_improved.publisher()', tech=True)
inp("Per-sensor overhead (+ sensor name length)", "sensor_over_b", 55, "bytes", '"<name>": {"val": 23.45, "status": "OK", "alarm": "NORMAL"}, — 53 B + 2 B separator', tech=True)
inp("Per-device overhead (+ device id length)", "device_over_b", 8, "bytes", '"<device>": {...}, — 6 B + 2 B separator', tech=True)
inp("System payload", "sys_payload_b", 320, "bytes", "get_system_status(): ts, gateway, city, city_time, platform, cpu/ram/disk, ip, os", tech=True)
inp("Fields added by the IoT rule", "iot_added_b", 46, "bytes", "SELECT *, clientid() AS gw, timestamp() AS rx_ms — stored in S3, read by Athena", tech=True)
inp("TLS record overhead", "tls_b", 29, "bytes/packet", "TLS 1.2 AES-GCM: 5 header + 8 nonce + 16 tag", tech=True)
inp("TCP/IP header overhead", "tcpip_b", 52, "bytes/packet", "IPv4 20 + TCP 20 + timestamps option 12", tech=True)
inp("Firehose buffer interval", "fh_buffer_s", 60, "seconds", "BufferingHints IntervalInSeconds in infra/template.yaml → one S3 object per stream per minute", tech=True)
inp("Athena minimum billed per query", "athena_min_mb", 10, "MB", "Athena bills at least 10 MB per query", tech=True)
inp("API calls per page load", "api_per_load", 5, "requests", "config, live, tests/active, tests list, reports list", tech=True)
inp("API calls per test", "api_per_test", 8, "requests", "start, stop, comments, report status polls", tech=True)
inp("DynamoDB writes per test", "ddb_writes_test", 10, "writes", "test item, ACTIVE state put/delete, comments, report status", tech=True)
inp("Dashboard files per page load", "web_files", 15, "files", "web/ has 14 files + index (Highcharts comes from its own CDN, not AWS)", tech=True)
inp("Dashboard size per page load", "web_kb", 47, "KB", "du -b web/ = 47,121 bytes (served uncompressed, CachingDisabled policy)", tech=True)
inp("/live response fixed part", "live_base_b", 1200, "bytes", "Estimate: timestamps, system block, JSON framing", tech=True)
inp("/live response per sensor", "live_sensor_b", 110, "bytes", "Estimate: key + val/status/alarm", tech=True)
inp("/tests/active response", "active_resp_b", 300, "bytes", "Estimate", tech=True)
inp("History/alarms response", "history_resp_kb", 50, "KB", "Estimate: ~2000 bucketed points", tech=True)
inp("API Lambda memory", "api_mem_mb", 512, "MB", "Globals.Function.MemorySize (arm64)", tech=True)
inp("API Lambda duration (simple call)", "api_ms", 100, "ms", "Estimate: S3/DynamoDB read", tech=True)
inp("API Lambda duration (Athena call)", "athena_ms", 4000, "ms", "Estimate: Lambda waits for the Athena query to finish", tech=True)
inp("Report Lambda memory", "report_mem_mb", 1024, "MB", "ReportFunction MemorySize", tech=True)
inp("Report Lambda duration", "report_s", 30, "seconds", "Estimate: 3 Athena queries + HTML/XLSX build", tech=True)
inp("Report size (HTML + XLSX)", "report_mb", 0.2, "MB", "Estimate", tech=True)
inp("Log volume per Lambda invocation", "log_kb", 0.5, "KB", "START/END/REPORT lines + app logs. No log retention is set, so logs accumulate", tech=True)
inp("Glue catalog requests per Athena query", "glue_per_query", 4, "requests", "Estimate: GetDatabase/GetTable/GetPartitions", tech=True)

dvf = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
ins.add_data_validation(dvf)
dvf.add(f"C{free_row}")
ins.cell(row=free_row, column=6, value=f'=IF(C{free_row}="Yes",1,0)').font = BLACK
name("free_flag", f"Inputs!$F${free_row}")
ins.freeze_panes = "A5"

# ---------------------------------------------------------------- Prices
ps = wb.create_sheet("Prices")
ps["A1"] = "AWS unit prices — us-east-1 (N. Virginia), on-demand, USD"
ps["A1"].font = TITLE
ps["A2"] = ("Prices are list prices as known when this file was built (Oct 2026) and are NOT fetched live. "
            "Check each against the source link before quoting; edit the blue cells to update every calculation.")
ps["A2"].font = NOTE
header(ps, 4, ["Service", "Item", "Price", "Per", "Always-free allowance (per account / month)", "Source"],
       [18, 40, 16, 26, 30, 60])
ps.row_dimensions[4].height = 30
prices = [
    ("IoT Core", "Connectivity", 0.08, "million connection-minutes", 0, "https://aws.amazon.com/iot-core/pricing/", "p_iot_conn"),
    ("IoT Core", "Messaging (5 KB increments)", 1.00, "million messages", 0, "https://aws.amazon.com/iot-core/pricing/", "p_iot_msg"),
    ("IoT Core", "Rules triggered", 0.15, "million rules", 0, "https://aws.amazon.com/iot-core/pricing/", "p_iot_rule"),
    ("IoT Core", "Rule actions executed", 0.15, "million actions", 0, "https://aws.amazon.com/iot-core/pricing/", "p_iot_action"),
    ("Data Firehose", "Direct PUT ingestion (5 KB record rounding)", 0.029, "GB ingested", 0, "https://aws.amazon.com/firehose/pricing/", "p_fh"),
    ("S3", "Standard storage", 0.023, "GB-month", 0, "https://aws.amazon.com/s3/pricing/", "p_s3_gb"),
    ("S3", "PUT / COPY / POST / LIST requests", 0.005, "1,000 requests", 0, "https://aws.amazon.com/s3/pricing/", "p_s3_put"),
    ("S3", "GET / SELECT requests", 0.0004, "1,000 requests", 0, "https://aws.amazon.com/s3/pricing/", "p_s3_get"),
    ("Athena", "Data scanned (10 MB min per query)", 5.00, "TB scanned", 0, "https://aws.amazon.com/athena/pricing/", "p_athena"),
    ("Glue", "Data Catalog requests", 1.00, "million requests", 1, "https://aws.amazon.com/glue/pricing/", "p_glue"),
    ("Lambda", "Requests", 0.20, "million requests", 1, "https://aws.amazon.com/lambda/pricing/", "p_lambda_req"),
    ("Lambda", "Compute, arm64", 0.0000133334, "GB-second", 400000, "https://aws.amazon.com/lambda/pricing/", "p_lambda_gbs"),
    ("API Gateway", "REST API requests", 3.50, "million requests", 0, "https://aws.amazon.com/api-gateway/pricing/", "p_apigw"),
    ("DynamoDB", "On-demand read request units", 0.125, "million RRU", 0, "https://aws.amazon.com/dynamodb/pricing/on-demand/", "p_ddb_r"),
    ("DynamoDB", "On-demand write request units", 0.625, "million WRU", 0, "https://aws.amazon.com/dynamodb/pricing/on-demand/", "p_ddb_w"),
    ("DynamoDB", "Storage", 0.25, "GB-month", 25, "https://aws.amazon.com/dynamodb/pricing/on-demand/", "p_ddb_gb"),
    ("CloudFront", "HTTPS requests (North America)", 1.00, "million requests", 10, "https://aws.amazon.com/cloudfront/pricing/", "p_cf_req"),
    ("CloudFront", "Data transfer out (North America)", 0.085, "GB", 1024, "https://aws.amazon.com/cloudfront/pricing/", "p_cf_gb"),
    ("Data transfer", "Out to internet (API responses)", 0.09, "GB", 100, "https://aws.amazon.com/ec2/pricing/on-demand/#Data_Transfer", "p_dto"),
    ("CloudWatch Logs", "Ingestion", 0.50, "GB", 5, "https://aws.amazon.com/cloudwatch/pricing/", "p_cw_in"),
    ("CloudWatch Logs", "Storage", 0.03, "GB-month", 5, "https://aws.amazon.com/cloudwatch/pricing/", "p_cw_gb"),
    ("EventBridge", "Scheduled rule (weekly report)", 0, "invocation", 0, "https://aws.amazon.com/eventbridge/pricing/", "p_eb"),
    ("IAM / API keys / usage plan", "No charge", 0, "-", 0, "-", "p_free"),
]
for i, (svc, item, price, per, free, src, nm) in enumerate(prices):
    row = 5 + i
    ps.cell(row=row, column=1, value=svc).font = BLACK
    ps.cell(row=row, column=2, value=item).font = BLACK
    c = ps.cell(row=row, column=3, value=price)
    c.font, c.border, c.fill = BLUE, BOX, YELLOW
    c.number_format = USDP if price and price < 0.001 else USD4
    ps.cell(row=row, column=4, value=per).font = BLACK
    fc = ps.cell(row=row, column=5, value=free)
    fc.font, fc.number_format = BLUE, NUM
    ps.cell(row=row, column=6, value=src).font = NOTE
    name(nm, f"Prices!$C${row}")
    name(nm + "_free", f"Prices!$E${row}")
pr = 5 + len(prices) + 1
for t in ["Not included: IoT Device Management/Defender (not used), Route 53 (no custom domain), taxes, support plans.",
          "IoT data sent INTO AWS has no transfer charge. S3 lifecycle rule raw/→Standard-IA after 90 days does not apply: since Sept 2024 S3 skips transitions of objects < 128 KB, and Firehose writes ~1 small object per minute, so raw data stays in Standard.",
          "AWS 12-month free tier / new-account credits are not modelled; only always-free allowances (toggle on the Inputs sheet)."]:
    ps.cell(row=pr, column=1, value=t).font = NOTE
    pr += 1

# ---------------------------------------------------------------- Calculation
cs = wb.create_sheet("Calculation")
cs["A1"] = "Traffic and usage (per site, per month unless noted)"
cs["A1"].font = TITLE
cs["A2"] = "All cells are formulas — edit Sensors, Inputs or Prices instead."
cs["A2"].font = NOTE
header(cs, 4, ["Quantity", "Name", "Value", "Unit", "How it is calculated"], [48, 26, 18, 22, 80])
r = 5


def calc(label, nm, formula, unit, note, fmt=NUM):
    global r
    cs.cell(row=r, column=1, value=label).font = BLACK
    cs.cell(row=r, column=2, value=nm).font = NOTE
    c = cs.cell(row=r, column=3, value=formula)
    c.font, c.number_format, c.border = BLACK, fmt, BOX
    cs.cell(row=r, column=4, value=unit).font = BLACK
    cs.cell(row=r, column=5, value=note).font = NOTE
    name(nm, f"Calculation!$C${r}")
    r += 1


def csec(text):
    global r
    r += 1
    section(cs, r, text, 5)
    r += 1


ATH_MIN = "athena_min_mb*1048576"
GB = "1073741824"
csec("Sensors and payloads")
calc("Active sensors", "active_sensors", f'=COUNTIFS(Sensors!$D${FIRST}:$D${LAST},1,Sensors!$C${FIRST}:$C${LAST},"<>")', "sensors", "Rows on Sensors with Active = 1")
calc("Devices with active sensors", "active_devices", f"=COUNTIF({SENS_F},\">0\")", "devices", "Distinct Device IDs")
calc("Data message payload", "data_payload_b", f"=data_base_b+SUM({SENS_E})+SUM({SENS_F})", "bytes", "Fixed part + every active sensor + every device wrapper")
calc("System message payload", "sys_payload", "=sys_payload_b", "bytes", "Independent of sensor count")
calc("Data messages per day", "data_msgs_day", "=86400/poll_s", "messages/day", "86,400 s ÷ data interval")
calc("System messages per day", "sys_msgs_day", "=86400/sys_s", "messages/day", "86,400 s ÷ system interval")

csec("Network traffic Pi ↔ AWS IoT (MQTT over TLS, QoS 1)")
calc("Data packet on the wire (upload)", "data_up_b", "=data_payload_b+6+LEN(topic_data)+tls_b+tcpip_b", "bytes", "Payload + MQTT header (fixed 2 + topic length 2 + topic + packet id 2) + TLS + TCP/IP")
calc("System packet on the wire (upload)", "sys_up_b", "=sys_payload+6+LEN(topic_sys)+tls_b+tcpip_b", "bytes", "Same, for the system topic")
calc("PUBACK received per message (download)", "ack_down_b", "=4+tls_b+tcpip_b", "bytes", "QoS 1 acknowledgement from AWS")
calc("TCP ACK sent per message (upload)", "ack_up_b", "=tcpip_b", "bytes", "Pi acknowledges the PUBACK segment")
calc("Payload per day", "payload_day_mb", "=(data_msgs_day*data_payload_b+sys_msgs_day*sys_payload)/1000000", "MB/day", "JSON only", NUM2)
calc("Total on the wire per day", "wire_day_mb", "=(data_msgs_day*(data_up_b+ack_down_b+ack_up_b)+sys_msgs_day*(sys_up_b+ack_down_b+ack_up_b))/1000000", "MB/day", "Upload + download, all overhead included (excludes rare TLS reconnects, ~7 KB each)", NUM2)
calc("Payload per month", "payload_month_mb", "=payload_day_mb*days_month", "MB/month", "", NUM2)
calc("Total on the wire per month", "wire_month_mb", "=wire_day_mb*days_month", "MB/month", "Use this to size the Pi's mobile data plan", NUM2)
calc("MQTT messages per month", "msgs_month", "=(data_msgs_day+sys_msgs_day)*days_month", "messages", "")
calc("IoT billed messages per month", "iot_billed_month", "=(data_msgs_day*ROUNDUP(data_payload_b/5120,0)+sys_msgs_day*ROUNDUP(sys_payload/5120,0))*days_month", "messages", "Each message billed per started 5 KB; above ~90 sensors a data message counts twice")

csec("Storage pipeline (IoT rules → Firehose → S3, plus latest/*.json)")
calc("Stored data record", "data_rec_b", "=data_payload_b+iot_added_b", "bytes", "Payload + gw + rx_ms added by the IoT rule")
calc("Stored system record", "sys_rec_b", "=sys_payload+iot_added_b", "bytes", "")
calc("Raw data bytes per day (data table)", "raw_day_data", "=data_msgs_day*data_rec_b", "bytes/day", "What Athena scans per day of history")
calc("Raw data bytes per day (system table)", "raw_day_sys", "=sys_msgs_day*sys_rec_b", "bytes/day", "")
calc("New raw data per month", "raw_gb_month", f"=(raw_day_data+raw_day_sys)*days_month/{GB}", "GB/month", "Uncompressed JSON (Firehose CompressionFormat UNCOMPRESSED)", NUM4)
calc("Firehose billed volume", "fh_gb_month", f"=(data_msgs_day*ROUNDUP(data_rec_b/5120,0)+sys_msgs_day*ROUNDUP(sys_rec_b/5120,0))*5120*days_month/{GB}", "GB/month", "Every record rounded up to 5 KB — small records cost far more than their size", NUM4)
calc("S3 objects per day, data stream", "objs_day_data", "=MIN(data_msgs_day,86400/fh_buffer_s)", "objects/day", "One object per buffer interval while data flows")
calc("S3 objects per day, system stream", "objs_day_sys", "=MIN(sys_msgs_day,86400/fh_buffer_s)", "objects/day", "")
calc("Firehose S3 PUTs per month", "fh_puts_month", "=(objs_day_data+objs_day_sys)*days_month", "requests", "")
calc("latest/*.json S3 PUTs per month", "latest_puts_month", "=msgs_month", "requests", "IoT rule S3 action overwrites latest/data.json or latest/system.json on EVERY message")
calc("S3 storage at retention horizon", "s3_store_gb", f"=raw_gb_month*months_kept+(tests_month+weekly_reports)*report_mb*months_kept/1024", "GB", "Raw history + reports after 'months kept'; grows linearly forever (no expiry)", NUM4)

csec("Dashboard and API")
calc("Live polls per month", "live_polls", "=viewers*hours_open*3600/live_poll_s*days_month", "polls", "Each poll = GET /live + GET /tests/active")
calc("Page loads per month", "page_loads_month", "=page_loads*days_month", "loads", "")
calc("History/alarm queries per month", "history_month", "=history_q*days_month", "queries", "")
calc("Reports per month", "reports_month", "=tests_month+weekly_reports", "reports", "Test reports + weekly reports")
calc("Simple API requests", "api_simple", "=live_polls*2+page_loads_month*api_per_load+tests_month*api_per_test", "requests", "No Athena involved")
calc("Athena-backed API requests", "api_athena", "=history_month+exports", "requests", "/history, /alarms, /export")
calc("API Gateway requests", "api_requests", "=api_simple+api_athena", "requests", "")
calc("Lambda invocations", "lambda_inv", "=api_requests+reports_month", "invocations", "API function + report function")
calc("Lambda compute", "lambda_gbs", "=(api_simple*api_ms/1000+api_athena*athena_ms/1000)*api_mem_mb/1024+reports_month*report_s*report_mem_mb/1024", "GB-seconds", "Duration × memory", NUM2)
calc("API response data out", "api_out_gb", f"=(live_polls*(live_base_b+live_sensor_b*active_sensors+active_resp_b)+api_athena*history_resp_kb*1024)/{GB}", "GB", "", NUM4)
calc("DynamoDB reads", "ddb_reads", "=live_polls+page_loads_month*2+tests_month*api_per_test", "read units", "/tests/active on every poll, test lists, comments")
calc("DynamoDB writes", "ddb_writes", "=tests_month*ddb_writes_test", "write units", "")
calc("CloudFront requests", "cf_requests", "=page_loads_month*web_files", "requests", "")
calc("CloudFront data out", "cf_gb", f"=page_loads_month*web_kb*1024/{GB}", "GB", "", NUM4)
calc("S3 GETs from dashboard", "s3_get_dash", "=live_polls*2+page_loads_month*web_files+hours_open*12*days_month", "requests", "latest/data + latest/system per poll; web files (CloudFront caching disabled); config.json refresh every 5 min")

csec("Athena (history, exports, reports)")
calc("Scan per history/export query", "q_hist_b", f"=MAX({ATH_MIN},raw_day_data*(history_days+1))", "bytes", "Whole dt= day partitions in the range (+1 boundary day), minimum 10 MB")
calc("Days scanned per test report", "test_days", "=ROUNDUP(test_hours/24,0)+1", "days", "")
calc("Scan per test report", "q_test_b", f"=2*MAX({ATH_MIN},raw_day_data*test_days)+MAX({ATH_MIN},raw_day_sys*test_days)", "bytes", "sensor_summary + data_gaps (data table) + system_summary")
calc("Scan per weekly report", "q_week_b", f"=2*MAX({ATH_MIN},raw_day_data*8)+MAX({ATH_MIN},raw_day_sys*8)", "bytes", "Same 3 queries over 7 days (+1 boundary)")
calc("Athena queries per month", "athena_queries", "=api_athena+3*reports_month", "queries", "")
calc("Athena data scanned", "athena_tb", "=(api_athena*q_hist_b+tests_month*q_test_b+weekly_reports*q_week_b)/1099511627776", "TB/month", "", '0.000000')
calc("S3 GETs by Athena", "s3_get_athena", "=api_athena*(history_days+1)*objs_day_data+tests_month*test_days*(2*objs_day_data+objs_day_sys)+weekly_reports*8*(2*objs_day_data+objs_day_sys)", "requests", "Athena reads every small Firehose object in the range; billed to the bucket owner")
calc("S3 PUTs other", "s3_put_other", "=athena_queries*2+exports+reports_month*2", "requests", "Athena results (csv + metadata), exports, report files")
calc("Glue catalog requests", "glue_requests", "=athena_queries*glue_per_query", "requests", "")
calc("CloudWatch log ingestion", "logs_gb", f"=lambda_inv*log_kb*1024/{GB}", "GB/month", "", NUM4)
calc("CloudWatch log storage at retention horizon", "logs_store_gb", "=logs_gb*months_kept", "GB", "No retention configured on the Lambda log groups", NUM4)

# ---- cost table
r += 2
cs.cell(row=r, column=1, value="Monthly cost per site (USD)").font = TITLE
r += 1
cost_hdr = r
for i, t in enumerate(["Service", "Item", "Usage", "Usage unit", "Free allowance", "Billable", "Unit price", "Monthly cost"], 1):
    c = cs.cell(row=r, column=i, value=t)
    c.font, c.fill, c.border = HDR, HDR_FILL, BOX
for col, w in zip("FGH", (16, 18, 16)):
    cs.column_dimensions[col].width = w
r += 1
costs = [
    ("IoT Core", "Connectivity (always connected)", "=1440*days_month/1000000", "million conn-min", "p_iot_conn"),
    ("IoT Core", "Messaging", "=iot_billed_month/1000000", "million messages", "p_iot_msg"),
    ("IoT Core", "Rules triggered", "=msgs_month/1000000", "million rules", "p_iot_rule"),
    ("IoT Core", "Rule actions (Firehose + S3 per message)", "=msgs_month*2/1000000", "million actions", "p_iot_action"),
    ("Data Firehose", "Ingestion (5 KB rounding)", "=fh_gb_month", "GB", "p_fh"),
    ("S3", "PUT: latest/*.json on every message", "=latest_puts_month/1000", "1,000 requests", "p_s3_put"),
    ("S3", "PUT: Firehose objects", "=fh_puts_month/1000", "1,000 requests", "p_s3_put"),
    ("S3", "PUT: Athena results, exports, reports", "=s3_put_other/1000", "1,000 requests", "p_s3_put"),
    ("S3", "GET: dashboard (live, web files, config)", "=s3_get_dash/1000", "1,000 requests", "p_s3_get"),
    ("S3", "GET: Athena reading raw objects", "=s3_get_athena/1000", "1,000 requests", "p_s3_get"),
    ("S3", "Storage (at retention horizon)", "=s3_store_gb", "GB-month", "p_s3_gb"),
    ("Athena", "Data scanned", "=athena_tb", "TB", "p_athena"),
    ("Glue", "Data Catalog requests", "=glue_requests/1000000", "million requests", "p_glue"),
    ("Lambda", "Requests", "=lambda_inv/1000000", "million requests", "p_lambda_req"),
    ("Lambda", "Compute (arm64)", "=lambda_gbs", "GB-seconds", "p_lambda_gbs"),
    ("API Gateway", "REST requests", "=api_requests/1000000", "million requests", "p_apigw"),
    ("DynamoDB", "Reads", "=ddb_reads/1000000", "million RRU", "p_ddb_r"),
    ("DynamoDB", "Writes", "=ddb_writes/1000000", "million WRU", "p_ddb_w"),
    ("DynamoDB", "Storage", "=0.001", "GB-month", "p_ddb_gb"),
    ("CloudFront", "HTTPS requests", "=cf_requests/1000000", "million requests", "p_cf_req"),
    ("CloudFront", "Data transfer out", "=cf_gb", "GB", "p_cf_gb"),
    ("Data transfer", "API responses to internet", "=api_out_gb", "GB", "p_dto"),
    ("CloudWatch Logs", "Ingestion", "=logs_gb", "GB", "p_cw_in"),
    ("CloudWatch Logs", "Storage (at retention horizon)", "=logs_store_gb", "GB-month", "p_cw_gb"),
    ("EventBridge", "Weekly schedule", "=weekly_reports", "invocations", "p_eb"),
    ("IAM / API keys / usage plan", "No charge", "=0", "-", "p_free"),
]
first_cost = r
for svc, item, qty, unit, p in costs:
    cs.cell(row=r, column=1, value=svc).font = BLACK
    cs.cell(row=r, column=2, value=item).font = BLACK
    cs.cell(row=r, column=3, value=qty).number_format = '#,##0.000000'
    cs.cell(row=r, column=4, value=unit).font = BLACK
    cs.cell(row=r, column=5, value=f"={p}_free*free_flag").font = GREEN
    cs.cell(row=r, column=5).number_format = NUM2
    cs.cell(row=r, column=6, value=f"=MAX(0,C{r}-E{r})").number_format = '#,##0.000000'
    pc = cs.cell(row=r, column=7, value=f"={p}")
    pc.font, pc.number_format = GREEN, USD4
    cs.cell(row=r, column=8, value=f"=F{r}*G{r}").number_format = USD4
    for col in range(1, 9):
        cs.cell(row=r, column=col).border = BOX
        if cs.cell(row=r, column=col).font != GREEN and col not in (1, 2, 4):
            cs.cell(row=r, column=col).font = BLACK
    r += 1
last_cost = r - 1
cs.cell(row=r, column=1, value="Total per site per month").font = BOLD
tc = cs.cell(row=r, column=8, value=f"=SUM(H{first_cost}:H{last_cost})")
tc.font, tc.number_format, tc.border = BOLD, USD, BOX
name("cost_site_month", f"Calculation!$H${r}")
cs.cell(row=r + 1, column=1, value=("Free allowances are per AWS account; with several sites the model applies them to each site, so leave the toggle on 'No' "
                                     "for a conservative multi-site figure.")).font = NOTE
COST_A = f"Calculation!$A${first_cost}:$A${last_cost}"
COST_H = f"Calculation!$H${first_cost}:$H${last_cost}"
cs.freeze_panes = "A5"

# ---------------------------------------------------------------- Data Plan
dp = wb.create_sheet("Data Plan")
dp["A1"] = "Cellular data plan per Pi (gateway → AWS IoT plus everything else the Pi does online)"
dp["A1"].font = TITLE
dp["A2"] = ("Blue = edit · black = formula · green = from another sheet. Sensor traffic comes from the Calculation sheet; "
            "the other lines are assumptions — measure the real usage on the Pi after a week (vnstat or /sys/class/net/<iface>/statistics).")
dp["A2"].font = NOTE
header(dp, 4, ["Item", "Name", "Value", "Unit", "Source / note"], [46, 22, 14, 18, 90])
r = 5


def dpl(label, nm, value, unit, note, kind="input", fmt=NUM2, key=False):
    global r
    dp.cell(row=r, column=1, value=label).font = BOLD if kind == "total" else BLACK
    dp.cell(row=r, column=2, value=nm).font = NOTE
    c = dp.cell(row=r, column=3, value=value)
    c.font = {"input": BLUE, "link": GREEN, "calc": BLACK, "total": BOLD}[kind]
    c.number_format, c.border = fmt, BOX
    if key:
        c.fill = YELLOW
    dp.cell(row=r, column=4, value=unit).font = BLACK
    dp.cell(row=r, column=5, value=note).font = NOTE
    name(nm, f"'Data Plan'!$C${r}")
    r += 1


def dsec(text):
    global r
    r += 1
    section(dp, r, text, 5)
    r += 1


dsec("Gateway traffic (sensor + system messages)")
dpl("Gateway MQTT traffic", "dp_gateway_mb", "=wire_month_mb", "MB/month",
    "Calculation sheet: payload + MQTT/TLS/TCP overhead + acks, for the active sensors and intervals", kind="link")

dsec("Other traffic on the Pi (assumptions)")
dpl("MQTT reconnects", "dp_reconnects", 5, "per day", "Weak cellular signal → more reconnects. Each one repeats the TLS handshake", key=True)
dpl("Traffic per reconnect", "dp_reconnect_kb", 7, "KB", "TLS handshake with certificate chain + MQTT CONNECT")
dpl("TCP retransmissions", "dp_retrans_pct", 0.1, "% of gateway traffic", "Typically 5–15 % on cellular", fmt="0.0%", key=True)
dpl("DNS + NTP time sync", "dp_dns_ntp_mb", 5, "MB/month", "Estimate")
dpl("Remote SSH / maintenance", "dp_ssh_mb", 10, "MB/month", "More if you copy logs or databases off the Pi")
dpl("Code/config updates (push_site.sh)", "dp_deploys", 2, "per month", "rsync of the Akvo_Green code + config.json")
dpl("Traffic per update", "dp_deploy_mb", 3, "MB", "Estimate; rsync only sends changed files")
dpl("OS updates (apt / unattended-upgrades)", "dp_os_mb", 0, "MB/month",
    "0 = automatic updates disabled and done over Wi-Fi/Ethernet. If left on, budget 300+ MB (a kernel/firmware update alone can exceed this)", key=True)
dpl("Carrier minimum billed per session", "dp_session_min_kb", 0, "KB",
    "0 = billed by actual usage. Some M2M plans round each data session up (e.g. 100 KB); every reconnect starts a new session", key=True)
dpl("Safety margin", "dp_margin", 0.3, "%", "Headroom for unexpected traffic and month-to-month variation", fmt="0%", key=True)

dsec("Monthly data needed per Pi")
dpl("Reconnects", "dp_reconnect_mb", "=dp_reconnects*dp_reconnect_kb*days_month/1000", "MB/month", "Reconnects × KB × days", kind="calc")
dpl("Retransmissions", "dp_retrans_mb", "=dp_gateway_mb*dp_retrans_pct", "MB/month", "Gateway traffic × %", kind="calc")
dpl("Updates", "dp_updates_mb", "=dp_deploys*dp_deploy_mb+dp_os_mb", "MB/month", "Code updates + OS updates", kind="calc")
dpl("Carrier session rounding", "dp_rounding_mb", "=dp_reconnects*days_month*MAX(0,dp_session_min_kb-dp_reconnect_kb)/1000",
    "MB/month", "Extra billed when each session is rounded up to the carrier minimum", kind="calc")
dpl("Subtotal", "dp_subtotal_mb",
    "=dp_gateway_mb+dp_reconnect_mb+dp_retrans_mb+dp_dns_ntp_mb+dp_ssh_mb+dp_updates_mb+dp_rounding_mb", "MB/month", "", kind="calc")
dpl("Data needed per Pi (with margin)", "dp_required_mb", "=dp_subtotal_mb*(1+dp_margin)", "MB/month", "Size the plan with this", kind="total")
dpl("Per day", "dp_required_day_mb", "=dp_required_mb/days_month", "MB/day", "", kind="calc")
dpl("All Pis (× sites)", "dp_required_all_gb", "=dp_required_mb*sites/1000", "GB/month", "If the SIMs share a pooled plan", kind="calc")

dsec("Plan options (list from smallest to largest; enter your carrier's plans and prices)")
dpl("Currency", "dp_currency", "MXN", "text", "Currency of the prices below")
r += 1
for i, t in enumerate(["Plan", "Data (MB/month)", "Price / month", "Fits?", "Headroom"], 1):
    c = dp.cell(row=r, column=i, value=t)
    c.font, c.fill, c.border = HDR, HDR_FILL, BOX
r += 1
plan_first = r
for label, mb in [("250 MB", 250), ("500 MB", 500), ("1 GB", 1000), ("2 GB", 2000), ("3 GB", 3000), ("5 GB", 5000)]:
    for col, v in ((1, label), (2, mb), (3, None)):
        c = dp.cell(row=r, column=col, value=v)
        c.font, c.border = BLUE, BOX
    dp.cell(row=r, column=2).number_format = NUM
    dp.cell(row=r, column=3).number_format = NUM2
    dp.cell(row=r, column=3).fill = YELLOW
    f = dp.cell(row=r, column=4, value=f'=IF(B{r}>=dp_required_mb,"Yes","No")')
    h = dp.cell(row=r, column=5, value=f"=IFERROR(B{r}/dp_required_mb-1,0)")
    f.font, h.font, h.number_format = BLACK, BLACK, '0%;-0%;"-"'
    f.border = h.border = BOX
    r += 1
plan_last = r - 1
PLAN_MB = f"'Data Plan'!$B${plan_first}:$B${plan_last}"
PLAN_NAME = f"'Data Plan'!$A${plan_first}:$A${plan_last}"
PLAN_PRICE = f"'Data Plan'!$C${plan_first}:$C${plan_last}"
dp.cell(row=r, column=1, value="Price column left empty = unknown; fill in your carrier's offer to see the monthly cost.").font = NOTE
r += 2
IDX = f'COUNTIF({PLAN_MB},"<"&dp_required_mb)+1'
FITS = f"dp_required_mb<=MAX({PLAN_MB})"
dpl("Recommended plan per Pi", "dp_plan", f'=IF({FITS},INDEX({PLAN_NAME},{IDX}),"Larger than the listed plans")',
    "", "Smallest listed plan that covers the data needed", kind="total", fmt="General")
dpl("Plan size", "dp_plan_mb", f"=IF({FITS},INDEX({PLAN_MB},{IDX}),0)", "MB/month", "", kind="calc", fmt=NUM)
dpl("Plan price per Pi", "dp_plan_price", f"=IF({FITS},INDEX({PLAN_PRICE},{IDX}),0)", "per month", "In the currency above", kind="calc")
dpl("Plan price, all Pis", "dp_plan_price_all", "=dp_plan_price*sites", "per month", "× number of sites", kind="calc")
dp.freeze_panes = "A5"

# ---------------------------------------------------------------- Summary
sm["A1"] = "VenkoDemo / Akvo_Green — AWS cost calculator"
sm["A1"].font = TITLE
sm["A2"] = ("How to use: edit sensors on 'Sensors' (Active 1/0), intervals and usage on 'Inputs' (yellow cells), unit prices on 'Prices'. "
            "Size the Pi's cellular plan on 'Data Plan'. Everything else recalculates. Region us-east-1, prices to be verified before quoting.")
sm["A2"].font = NOTE
sm.column_dimensions["A"].width = 44
sm.column_dimensions["B"].width = 18
sm.column_dimensions["C"].width = 16
sm.column_dimensions["D"].width = 60

rows = [
    ("Key inputs", None, None, None),
    ("Active sensors", "=active_sensors", NUM, "From the Sensors sheet"),
    ("Sensor data interval (s)", "=poll_s", NUM, "Inputs"),
    ("System data interval (s)", "=sys_s", NUM, "Inputs"),
    ("Sites", "=sites", NUM, "Inputs"),
    ("Data volume per site", None, None, None),
    ("Data message size (bytes)", "=data_payload_b", NUM, "JSON payload with all active sensors"),
    ("MQTT messages per month", "=msgs_month", NUM, "Data + system"),
    ("Payload per month (MB)", "=payload_month_mb", NUM2, "JSON only"),
    ("Network traffic per month (MB)", "=wire_month_mb", NUM2, "Gateway MQTT only, incl. MQTT/TLS/TCP overhead and acks"),
    ("New S3 raw data per month (GB)", "=raw_gb_month", NUM4, ""),
    ("Cellular data plan per Pi", None, None, None),
    ("Data needed per Pi (MB/month)", "=dp_required_mb", NUM2, "Data Plan sheet: gateway + other Pi traffic + margin"),
    ("Recommended plan", "=dp_plan", "General", "Smallest plan on the Data Plan sheet that covers it"),
    ("Cost", None, None, None),
    ("Monthly cost per site", "=cost_site_month", USD, "Calculation sheet, at the 'months kept' storage horizon"),
    ("Monthly cost, all sites", "=cost_site_month*sites", USD, ""),
    ("Yearly cost, all sites", "=cost_site_month*sites*12", USD, "Storage grows each month; this uses the horizon month ×12 (upper bound)"),
    ("Cost per active sensor per month", "=IFERROR(cost_site_month/active_sensors,0)", USD4, "Includes the fixed per-site costs spread over the sensors"),
]
r = 4
for label, f, fmt, note in rows:
    if f is None:
        section(sm, r, label, 4)
    else:
        sm.cell(row=r, column=1, value=label).font = BLACK
        c = sm.cell(row=r, column=2, value=f)
        c.font, c.number_format, c.border = GREEN, fmt, BOX
        sm.cell(row=r, column=4, value=note).font = NOTE
        if label.startswith("Monthly cost per site"):
            c.font = Font(name=F, bold=True, color="008000")
            c.fill = YELLOW
    r += 1

r += 1
section(sm, r, "Monthly cost per site by service", 4)
r += 1
for i, t in enumerate(["Service", "USD / month", "Share"], 1):
    c = sm.cell(row=r, column=i, value=t)
    c.font, c.fill, c.border = HDR, HDR_FILL, BOX
r += 1
svc_first = r
services = []
for svc, *_ in costs:
    if svc not in services:
        services.append(svc)
for svc in services:
    sm.cell(row=r, column=1, value=svc).font = BLACK
    c = sm.cell(row=r, column=2, value=f'=SUMIF({COST_A},A{r},{COST_H})')
    c.number_format, c.font, c.border = USD4, GREEN, BOX
    s = sm.cell(row=r, column=3, value=f"=IFERROR(B{r}/cost_site_month,0)")
    s.number_format, s.font, s.border = '0.0%;(0.0%);"-"', BLACK, BOX
    r += 1
svc_last = r - 1
sm.cell(row=r, column=1, value="Total").font = BOLD
t = sm.cell(row=r, column=2, value=f"=SUM(B{svc_first}:B{svc_last})")
t.number_format, t.font, t.border = USD, BOLD, BOX

ch = BarChart()
ch.type = "bar"
ch.title = "Monthly cost per site by service (USD)"
ch.style = 10
ch.legend = None
ch.add_data(Reference(sm, min_col=2, min_row=svc_first - 1, max_row=svc_last), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=svc_first, max_row=svc_last))
ch.height, ch.width = 9, 16
sm.add_chart(ch, f"F{svc_first - 2}")

# ---------------------------------------------------------------- How To (first tab)
ht = wb.create_sheet("How To", 0)
ht.column_dimensions["A"].width = 6
ht.column_dimensions["B"].width = 110
ht.column_dimensions["C"].width = 24
ht["A1"] = "How to use this calculator"
ht["A1"].font = TITLE
ht["A2"] = ("Only edit BLUE cells (yellow fill = the main settings). Black cells are formulas, green cells pull values from another sheet. "
            "Click a link in column C to jump to the cell.")
ht["A2"].font = NOTE


def ref(nm):
    """'Inputs!$C$7' -> (sheet, 'C7') for the defined name."""
    sheet, cell = wb.defined_names[nm].attr_text.rsplit("!", 1)
    return sheet.strip("'"), cell.replace("$", "")


r = 4


def task(title, steps):
    global r
    section(ht, r, title, 3)
    r += 1
    for i, step in enumerate(steps, 1):
        text, link = step if isinstance(step, tuple) else (step, None)
        ht.cell(row=r, column=1, value=i).font = BLACK
        ht.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="center")
        b = ht.cell(row=r, column=2, value=text)
        b.font, b.alignment = BLACK, Alignment(wrap_text=True, vertical="top")
        if link:
            if link in wb.sheetnames:
                sheet, cell = link, "A1"
            elif "!" in link:
                sheet, cell = link.split("!")
            else:
                sheet, cell = ref(link)
            # HYPERLINK() rather than a cell hyperlink: LibreOffice URL-encodes the space in
            # "Data Plan" when it re-saves cell hyperlinks, which Excel can't follow.
            c = ht.cell(row=r, column=3, value=f'=HYPERLINK("#\'{sheet}\'!{cell}","→ {sheet}!{cell}")')
            c.font = Font(name=F, color="0563C1", underline="single")
            c.alignment = Alignment(vertical="top")
        r += 1
    r += 1


task("1. Add a sensor", [
    ("Go to the Sensors sheet and find the first empty row (rows 6–205).", f"Sensors!B{FIRST + len(sensors)}"),
    "Type the Device ID exactly as in Akvo_Green config.json → devices[].id (for example DEV_16). Upper/lower case and length matter: the name is sent in every message.",
    "Type the Sensor name exactly as in config.json → devices[].sensors[].name (for example FlowSensor).",
    "Type 1 in Active. Columns E and F fill in by themselves.",
    "Several sensors on the same device: repeat the same Device ID on each row; the device is only counted once.",
    ("Check the result: Summary → Active sensors, Data message size, Monthly cost per site.", "Summary"),
])
task("2. Remove or pause a sensor", [
    "Set Active to 0 to leave the sensor out but keep the row (useful for what-if comparisons), or delete the Device ID and Sensor name.",
    "Do not delete whole rows or columns: the formulas expect the list in rows 6–205.",
])
task("3. Change the sampling frequency", [
    ("Sensor data interval: seconds between data messages (gateway.poll_interval in config.json). One message carries ALL sensors.", "poll_s"),
    ("System data interval: seconds between CPU/RAM/disk messages (gateway.system_interval).", "sys_s"),
    "Example: 20 → 10 s doubles the data messages, IoT messages, rule actions and S3 writes; 20 → 60 s cuts them to a third.",
    ("Keep the dashboard poll in line: if the data interval is longer than the live poll, the dashboard polls for data that is not there yet.", "live_poll_s"),
])
task("4. Model several sites (several Pis)", [
    ("Set Number of sites. Every site is its own stack (bucket, API, dashboard) and its own SIM, so AWS cost and data plan are multiplied.", "sites"),
    "The Sensors sheet describes ONE site. If the sites have very different sensors, make a copy of the file per site.",
])
task("5. Describe how the dashboard is used", [
    ("Dashboards open at the same time and Hours per day open drive API Gateway and Lambda (each open dashboard polls every 10 s).", "viewers"),
    ("A control-room screen open all day: Hours per day open = 24.", "hours_open"),
    ("History / alarm queries per day and their average range in days drive Athena and S3 GET cost.", "history_q"),
    ("Tests per month and their duration drive report generation.", "tests_month"),
])
task("6. Update AWS prices", [
    ("Open the Prices sheet and compare each blue price with the AWS page in the Source column (region us-east-1).", "Prices"),
    "Type the new price over the old one; keep the same unit as the 'Per' column (e.g. price per million requests, not per request).",
    ("Optional: Apply AWS always-free tier = Yes subtracts the always-free allowances. They are per AWS account, so with several sites leave it on No for a safe estimate.", "free_tier"),
])
task("7. Size the Pi's cellular data plan", [
    ("Open the Data Plan sheet. Gateway traffic is taken from your sensors and intervals automatically.", "dp_gateway_mb"),
    ("OS updates: 0 if automatic updates are disabled on the Pi; about 300 MB/month if they stay on.", "dp_os_mb"),
    ("MQTT reconnects per day: raise it if the cellular signal is weak (check the gateway log for 'Connecting MQTT').", "dp_reconnects"),
    ("Carrier minimum per session: ask the carrier whether data sessions are rounded up (M2M/IoT SIMs sometimes are). 0 = billed by actual usage.", "dp_session_min_kb"),
    ("Enter the carrier's plans in MB from smallest to largest, with their monthly price. The recommended plan is the smallest one that fits.", f"Data Plan!C{plan_first}"),
    ("Read the result: Data needed per Pi and Recommended plan per Pi.", "dp_plan"),
])
task("8. Check the estimate against reality", [
    "Data used by the Pi: on the Pi run   sudo apt install vnstat   and after a few days   vnstat -m   (or read /sys/class/net/<interface>/statistics/rx_bytes and tx_bytes). Compare with Data needed per Pi before the margin (Subtotal).",
    "Messages really sent: AWS console → IoT Core → Monitor, or CloudWatch metric AWS/IoT PublishIn.Success. Compare with Calculation → MQTT messages per month.",
    "Real AWS cost: AWS console → Billing → Cost Explorer, group by Service, last full month. Compare service by service with the table on the Summary sheet.",
    "If something differs a lot, adjust the matching grey technical constant on the Inputs sheet (for example API Lambda duration or log volume).",
])
task("9. Read the results", [
    ("Summary → Monthly cost per site (yellow) is the main figure; the table and chart show which services cost the most.", "Summary"),
    ("Calculation shows every intermediate quantity with how it is calculated, and the full cost table (usage → free allowance → billable → price → cost).", "Calculation"),
    "Storage figures are for the month after 'Months of history kept', because raw data in S3 and Lambda logs are never deleted.",
])
task("10. Rebuild the file from scratch", [
    "The workbook is generated by tools/build_cost_calculator.py in the VenkoDemo repo:   python3 tools/build_cost_calculator.py docs/aws_cost_calculator.xlsx",
    "Rebuilding overwrites every change made in the file (sensors, inputs, prices). Change the defaults in the script if you want them to survive a rebuild.",
    "Open the rebuilt file in Excel or LibreOffice once and save it, so the calculated values are stored.",
])
ht.freeze_panes = "A4"
wb.active = 0
sm.sheet_view.tabSelected = False

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
wb.save(OUT)
print("saved", OUT)
