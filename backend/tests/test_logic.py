import io
import json

import pytest
from openpyxl import load_workbook

from common import derived, queries, report_builder, timeutil

T0 = timeutil.parse_ms("2026-09-21T06:00:00Z")  # Monday 00:00 America/Mexico_City
POLL = 20_000

CONFIG = {
    "site": {"title": "Demo", "brand": "Ghotrix"},
    "gateway": {"gateway_id": "GW", "city": "America/Mexico_City", "poll_interval": 20},
    "sensors": [
        {"key": "DEV_5.PhSensor", "label": "pH", "unit": "pH", "min": 0, "max": 14},
        {"key": "DEV_1.Pt100Sensor", "label": "Temp <b>", "unit": "°C", "min": 0, "max": 55},
        {"key": "DEV_9.Silent", "label": "Sin datos", "unit": "x", "min": 0, "max": 1},
    ],
}


# ---------------------------------------------------------------- queries
def test_sensor_keys_validated():
    assert queries.parse_sensor_keys("DEV_5.PhSensor, DEV_1.Pt100Sensor") == ["DEV_5.PhSensor", "DEV_1.Pt100Sensor"]
    assert queries.parse_sensor_keys("") == []
    with pytest.raises(ValueError):
        queries.parse_sensor_keys("DEV_5.PhSensor' OR '1'='1")


def test_history_sql_filters_partition_time_and_sensors():
    sql = queries.history(T0, T0 + 3_600_000, ["DEV_5.PhSensor"], 60_000)
    assert "dt BETWEEN '2026-09-20' AND '2026-09-22'" in sql
    assert f"rx_ms >= {T0}" in sql and "IN ('DEV_5.PhSensor')" in sql
    assert "(rx_ms / 60000) * 60000" in sql


def test_system_stats_sql_covers_three_resources_in_range():
    sql = queries.system_stats(T0, T0 + 86_400_000)
    assert "FROM system_raw" in sql and f"rx_ms >= {T0}" in sql
    for c in ("cpu_load_percent", "ram_usage_percent", "disk_usage_percent"):
        assert f"min({c})" in sql and f"avg({c})" in sql and f"max({c})" in sql


# ------------------------------------------------------------------- time
def test_week_bounds_are_local_monday_to_monday():
    label, start, end = timeutil.week_bounds("2026-W39", "America/Mexico_City")
    assert start == T0 and end - start == 7 * 86_400_000
    assert timeutil.previous_iso_week(T0 + 8 * 86_400_000, "America/Mexico_City")[0] == "2026-W39"


def test_range_limits():
    with pytest.raises(ValueError):
        timeutil.validate_range(T0, T0)
    with pytest.raises(ValueError):
        timeutil.validate_range(T0, T0 + 40 * 86_400_000)
    assert timeutil.pick_bucket_ms(T0, T0 + 3_600_000) == 20_000
    assert timeutil.pick_bucket_ms(T0, T0 + 7 * 86_400_000) == 1_800_000


# --------------------------------------------------------------- episodes
def test_collapse_episodes_splits_on_gap_and_kind():
    rows = (
        [{"k": "A", "t": T0 + i * POLL, "val": 15 + i, "kind": "HIGH"} for i in range(5)]      # one episode
        + [{"k": "A", "t": T0 + (20 + i) * POLL, "val": 16, "kind": "HIGH"} for i in range(2)]  # after a gap
        + [{"k": "A", "t": T0 + 30 * POLL, "val": None, "kind": "BUS_ERROR"}]
        + [{"k": "B", "t": T0 + i * POLL, "val": -1 - i, "kind": "LOW"} for i in range(3)]
    )
    eps = report_builder.collapse_episodes(rows, max_gap_ms=int(POLL * 2.5))
    summary = sorted((e["k"], e["kind"], e["samples"], e["peak"]) for e in eps)
    assert summary == [("A", "BUS_ERROR", 1, None), ("A", "HIGH", 2, 16), ("A", "HIGH", 5, 19), ("B", "LOW", 3, -3)]


# ----------------------------------------------------------------- report
def _ctx():
    end = T0 + 3_600_000
    summary = [
        {"k": "DEV_5.PhSensor", "samples": 180, "ok": 180, "min": 6.9, "avg": 7.1, "max": 14.8, "std": 0.4,
         "first": 7.0, "last": 7.2, "low": 0, "high": 3},
        {"k": "DEV_1.Pt100Sensor", "samples": 180, "ok": 170, "min": 20, "avg": 22.5, "max": 25, "std": 1.1,
         "first": 21, "last": 24, "low": 0, "high": 0},
    ]
    series = [{"k": "DEV_5.PhSensor", "t": T0 + i * 60_000, "avg": 7 + i / 100, "min": 6.9, "max": 7.3, "n": 3}
              for i in range(60)]
    abnormal = [{"k": "DEV_5.PhSensor", "t": T0 + (100 + i) * POLL, "val": 14.5 + i / 10, "kind": "HIGH"} for i in range(3)]
    return report_builder.build_context(
        kind="test", title="Prueba </script><script>alert(1)</script>", config=CONFIG, from_ms=T0, to_ms=end,
        summary=summary, series=series, bucket_ms=60_000, abnormal=abnormal,
        gaps=[{"start": T0 + 600_000, "stop": T0 + 900_000}], system={"CPU": "10 / 20"},
        comments=[{"ts_ms": T0 + 120_000, "author": "op", "sensor": "DEV_5.PhSensor", "text": "ajuste <válvula>"}],
        meta={"operador": "op", "notas": ""}, generated_ms=end,
    )


def test_context_stats():
    ctx = _ctx()
    by = {s["key"]: s for s in ctx["sensors"]}
    assert by["DEV_5.PhSensor"]["status"] == "ALARMA" and by["DEV_5.PhSensor"]["alarm_pct"] == pytest.approx(1.7)
    assert by["DEV_1.Pt100Sensor"]["status"] == "OK" and by["DEV_1.Pt100Sensor"]["availability"] == pytest.approx(94.4)
    assert by["DEV_9.Silent"]["status"] == "SIN DATOS"
    assert ctx["alarm_sensor_count"] == 1
    assert ctx["episodes"][0]["peak"] == 14.7 and ctx["episodes"][0]["minutes"] == 1.0
    assert ctx["gaps"][0]["minutes"] == 5.0 and ctx["data_availability"] == pytest.approx(92.2)
    assert ctx["from_local"] == "2026-09-21 00:00:00"
    assert ctx["comments"][0]["sensor"] == "pH"


def test_html_escapes_user_text():
    html = report_builder.render_html(_ctx()).decode()
    assert "<script>alert(1)" not in html
    assert "ajuste &lt;válvula&gt;" in html
    assert "Temp <b>" not in html.split("const CHARTS")[0]
    charts = json.loads(html.split("const CHARTS = ")[1].split(";\n")[0].replace("<\\/", "</"))
    assert len(charts) == 3 and len(charts[0]["data"]) == 60


def test_xlsx_sheets():
    wb = load_workbook(io.BytesIO(report_builder.render_xlsx(_ctx())))
    assert wb.sheetnames == ["Resumen", "Lecturas (1.0 min)", "Alarmas", "Huecos de datos", "Comentarios", "Sistema"]
    readings = wb["Lecturas (1.0 min)"]
    assert readings.max_row == 61 and readings["A2"].value == "2026-09-21 00:00:00"
    assert wb["Alarmas"]["B2"].value == "HIGH"


# ------------------------------------------------------------ derived
REC = {"key": "REC", "label": "Recuperación", "unit": "%", "op": "ratio", "a": "F2", "b": "F1", "scale": 100,
       "min": 50, "max": 85, "decimals": 0}
BAL = {"key": "BAL", "label": "Balance", "unit": "gpm", "op": "balance", "a": "F1", "minus": ["F2", "F3"],
       "min": -1, "max": 1}


def test_derived_fixed_operations():
    vals = {"F1": 12.0, "F2": 8.4, "F3": 3.6, "P1": 45, "P2": 42}
    assert derived.value(REC, vals) == pytest.approx(70.0)
    assert derived.value(BAL, vals) == pytest.approx(0.0)
    assert derived.value({"op": "diff", "a": "P1", "b": "P2"}, vals) == 3
    assert derived.value(REC, {"F2": 8.4, "F1": 0}) is None          # no division by zero
    assert derived.value(BAL, {"F1": 12, "F2": 8}) is None           # a missing sensor -> no value


def test_derived_state_uses_warning_band():
    assert derived.state(REC, 70) == "ok"
    assert derived.state(REC, 83) == "warn"      # within 10 % of the 50-85 range below the max
    assert derived.state(REC, 90) == "alarm"
    assert derived.state(REC, None) == "none"


def test_report_context_includes_indicators_from_period_averages():
    cfg = {**CONFIG, "derived": [{**REC, "a": "DEV_1.Pt100Sensor", "b": "DEV_5.PhSensor", "provisional": True}]}
    ctx = report_builder.build_context(
        kind="weekly", title="S", config=cfg, from_ms=T0, to_ms=T0 + 3_600_000,
        summary=[{"k": "DEV_5.PhSensor", "ok": 1, "avg": 30.0}, {"k": "DEV_1.Pt100Sensor", "ok": 1, "avg": 21.0}],
        series=[], bucket_ms=60_000, abnormal=[], gaps=[], system=None, comments=[], generated_ms=T0)
    assert ctx["derived"] == [{"key": "REC", "label": "Recuperación", "unit": "%", "avg": 70.0, "lim_min": 50,
                               "lim_max": 85, "status": "OK", "provisional": True}]
    assert b"Indicadores del proceso" in report_builder.render_html(ctx)
    assert "Indicadores" in load_workbook(io.BytesIO(report_builder.render_xlsx(ctx))).sheetnames


# ------------------------------------------------------- warnings / summaries
def test_warn_bands_skip_bad_keys_and_missing_limits():
    sensors = [{"key": "DEV_3.P3", "min": 0, "max": 200},
               {"key": "DEV_1.T", "min": 0, "max": 50, "warn_margin": 0.2},
               {"key": "DEV_9.x' OR 1=1 --", "min": 0, "max": 1},
               {"key": "DEV_8.NoLimits", "min": None, "max": 10}]
    assert queries.warn_bands(sensors) == [("DEV_3.P3", 20.0, 180.0), ("DEV_1.T", 10.0, 40.0)]


def test_abnormal_sql_adds_warning_kinds_only_with_bands():
    plain = queries.abnormal_rows(T0, T0 + 3_600_000)
    assert "WARN_HIGH" not in plain and "WHERE kind IS NOT NULL" in plain
    sql = queries.abnormal_rows(T0, T0 + 3_600_000, bands=[("DEV_3.P3", 20.0, 180.0)])
    assert "= 'DEV_3.P3' AND r.val > 180.0) THEN 'WARN_HIGH'" in sql
    assert "= 'DEV_3.P3' AND r.val < 20.0) THEN 'WARN_LOW'" in sql


def test_range_stats_counts_alarm_warning_and_errors():
    sql = queries.range_stats(T0, T0 + 3_600_000, ["DEV_3.P3"], [("DEV_3.P3", 20.0, 180.0)])
    assert "count_if(kind IN ('WARN_LOW', 'WARN_HIGH')) AS warn" in sql and "IN ('DEV_3.P3')" in sql
    assert "max_by(val, t) FILTER (WHERE status = 'OK') AS last" in sql


def test_warning_episode_peak_and_summary():
    rows = [{"k": "A.p", "t": T0 + i * POLL, "val": 181 + i, "kind": "WARN_HIGH"} for i in range(3)]
    rows += [{"k": "A.p", "t": T0 + 86_400_000 + i * POLL, "val": 15 - i, "kind": "WARN_LOW"} for i in range(2)]
    rows += [{"k": "B.q", "t": T0 + 3_600_000, "val": None, "kind": "BUS_ERROR"},
             {"k": "B.q", "t": T0 + 7_200_000, "val": 210, "kind": "HIGH"}]
    eps = report_builder.collapse_episodes(rows, max_gap_ms=int(POLL * 2.5))
    peaks = {e["kind"]: e["peak"] for e in eps}
    assert peaks["WARN_HIGH"] == 183 and peaks["WARN_LOW"] == 14
    s = report_builder.alarm_summary(eps, {"A.p": "Presión"}, T0, T0 + 2 * 86_400_000, "America/Mexico_City")
    assert s["totals"] == {"alarm": 1, "warn": 2, "err": 1}
    assert s["top"] == {"key": "A.p", "label": "Presión", "count": 2}
    assert [d["day"] for d in s["days"]] == ["2026-09-21", "2026-09-22"]
    assert s["days"][0] == {"day": "2026-09-21", "alarm": 1, "warn": 1, "err": 1}
